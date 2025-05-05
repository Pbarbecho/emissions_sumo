import pandas as pd
import sys, os
from tqdm import tqdm
from joblib import Parallel, delayed, parallel_backend
import math
import time
import multiprocessing
import xml.etree.ElementTree as ET
import psutil
import shutil
import numpy as np
from datetime import datetime
import openpyxl


def time_diff(from_time, to_time):
    from_time_float = float(from_time)
    to_time_float = float(to_time)

    # Calculate the time difference
    if from_time_float == to_time_float:
        time_difference = 24.0
    elif to_time_float < from_time_float:
        to_time_float += 24
        time_difference = to_time_float - from_time_float
    else:
        time_difference = to_time_float - from_time_float

    time_difference_hours = int(time_difference)  # Número de horas completas entre from_time y to_time
    return time_difference_hours
    #return f"{time_difference_hours:05.2f}"


def generate_output_file(percentage_files_xls, output_file, factor):

    input_excel_list = percentage_files_xls['xls_traffic_file'].tolist()
    percentage_list = percentage_files_xls['porcentaje'].tolist()

    list_OD_files = []
    for k, file in enumerate(input_excel_list):

        # Load the Excel file
        df = pd.read_excel(file, index_col=0)

        OD_output_name = f"{output_file}/{k}.od"

        percentage_val = percentage_list[k]
        print(percentage_val, file)
        with open(OD_output_name, 'w') as f:
            # Write the header
            f.write('$OR;D2\n')
            f.write(f'* From-Time  To-Time\n{k}.00 {k+1}.00\n')
            f.write('* Factor\n')
            f.write(f'{factor}\n')
            f.write('* Porcentaje\n')
            f.write(f'* {percentage_val}\n')

            # Write the data
            for from_location, row in df.iterrows():
                for to_location, num_vehicles in row.items():
                    # adjusted_vehicles = num_vehicles * porcentaje
                    adjusted_vehicles = num_vehicles
                    # f.write(f'\t\t{from_location}          {to_location}       {num_vehicles}\n')
                    # f.write(f'\t\t\t{from_location}\t\t\t{to_location}\t\t\t{num_vehicles}\n')
                    #                f.write(f'\t\t{from_location.ljust(30)}{to_location.ljust(30)}{num_vehicles}\n')
                    f.write(f'\t\t{from_location.ljust(30)}{to_location.ljust(30)}{int(adjusted_vehicles)}\n')
        list_OD_files.append(OD_output_name)
    return list_OD_files


def gen_per_hour_traffic(porcentaje, archivo_entrada, archivo_salida, n, iter):
    #Cargar el archivo de entrada
    wb_entrada = openpyxl.load_workbook(archivo_entrada)
    wb_salida = openpyxl.Workbook()
    wb_salida.remove(wb_salida.active)  # Quita la hoja por defecto

    for hoja in wb_entrada.sheetnames:
        ws_entrada = wb_entrada[hoja]
        ws_salida = wb_salida.create_sheet(title=hoja)

        for fila in ws_entrada.iter_rows():
            nueva_fila = []
            for celda in fila:
                valor = celda.value
                if isinstance(valor, (int, float)):
                    nuevo_valor = round(valor * porcentaje)
                else:
                    nuevo_valor = valor  # Mantiene texto o celdas vacías igual
                nueva_fila.append(nuevo_valor)
            ws_salida.append(nueva_fila)
    # Guardar el nuevo archivo
    out_file_name = f"{archivo_salida}/{iter}.xlsx"
    wb_salida.save(out_file_name)
    return out_file_name


def create_files_per_hour(traffic_weights, archivo_entrada, archivo_salida, hours):
    list_out_files = []
    # falta validar probabilidad por caDA HORA para el trafico en forma de camello

    weights_list = peso_per_hour(traffic_weights)
    porcentajes_dict = dict(weights_list)

    # Lista para almacenar los datos
    datos = []

    for k in range(0, hours):
        porcentaje = porcentajes_dict[k]
        out_file_name = gen_per_hour_traffic(porcentaje, archivo_entrada, archivo_salida, hours, k)
        list_out_files.append(out_file_name)
        datos.append([porcentaje, out_file_name])
    # Crear el DataFrame
    df_result = pd.DataFrame(datos, columns=['porcentaje', 'xls_traffic_file'])
    return df_result


def peso_per_hour(pesos_hour_traffic_file):
    # Leer el archivo Excel de porcentajes
    df = pd.read_excel(pesos_hour_traffic_file)  # asegúrate de usar el nombre correcto

    # Suponemos que hay una columna llamada 'porcentaje'
    porcentajes = df["porcentaje"].astype(float).tolist()

    # Crear lista con total_hours igual al número de porcentajes
    total_hours = len(porcentajes)
    resultado = [(i, porcentajes[i]) for i in range(total_hours)]
    return resultado


def gen_od2trips_cfg(O_list, taz_file, out_dir):

    # read O files
    O_files_list = [os.path.basename(path.rstrip("/")) for path in O_list]
    O_listToStr = ','.join([f'{elem}' for elem in O_files_list])
    TAZ = taz_file

    od2trips_conf = os.path.join('templates', 'od2trips.cfg.xml')

    # Open original file
    tree = ET.parse(od2trips_conf)

    # Update O input
    parent = tree.find('input')
    ET.SubElement(parent, 'od-matrix-files').set('value', f'{O_listToStr}')
    ET.SubElement(parent, 'taz-files').set('value', f'{TAZ}')

    # Update output
    parent = tree.find('output')
    output_name = 'od2.trip.xml'
    ET.SubElement(parent, 'output-file').set('value', output_name)

    # Update seed number
    parent = tree.find('random_number')
    ET.SubElement(parent, 'seed').set('value', f'{1}')

    # Write xml
    cfg_name = f'{out_dir}/trips.cfg.xml'
    tree.write(cfg_name)
    return cfg_name, output_name


def exec_od2trips(fname):
    print('\nOD2Trips running .............\n')
    cmd = f'od2trips -v -c {fname}'
    os.system(cmd)
    # remove fromtotaz
    #output_file = f'{tripfile}.xml'
    #rm_taz = f"sed 's/fromTaz=\"{folders.O_district}\" toTaz=\"{folders.D_district}\"//' {tripfile} > {output_file}"
    #os.system(rm_taz)
    #return output_file


def gen_routes(net_file , cfg_name , trips_file, dua_output, i, hours, reroute_prob, path_sumo_outputs, edgeData_xml):
    #Generate configuration files for dua / ma router

    # Execute od2trips
    exec_od2trips(cfg_name)

    # Generate DUArouter cfg
    cfg_name, output_name = gen_DUArouter(net_file, trips_file, dua_output,  i, hours)

    # Generate sumo cfg
    output_name_m = "rutas_m.rou.xml"
    gen_sumo_cfg(output_name_m, dua_output, reroute_prob, hours, path_sumo_outputs)


def gen_sumo_cfg(routing_file, folders, rr_prob, hours, path_sumo_outputs):
    sumo_cfg = os.path.join('templates', 'osm.sumo.cfg')
    vtype = os.path.join('vtypes_m.add.xml')
    # new_emissions = os.path.join(folders.parents_dir,'templates', 'emissions.add.xml')
    TAZ = os.path.join('districts.taz.xml')
    net_file = 'network.net.xml'

    # Open original file
    tree = ET.parse(sumo_cfg)

    # Update rou input
    parent = tree.find('input')
    ET.SubElement(parent, 'net-file').set('value', f'{net_file}')
    ET.SubElement(parent, 'route-files').set('value', f'{routing_file}')

    edgeData = 'edgeData_cfg.xml'
    add_list = [vtype, edgeData]
    additionals = ','.join([elem for elem in add_list])

    # Update detector
    ET.SubElement(parent, 'additional-files').set('value', f'{additionals}')

    # Routing
    parent = tree.find('routing')
    ET.SubElement(parent, 'device.rerouting.probability').set('value', f'{rr_prob}')
    #ET.SubElement(parent, 'device.rerouting.output').set('value', f'{os.path.join(folders.reroute, "reroute.xml")}')

    # Update outputs
    parent = tree.find('output')

    # outputs
    outputs = ['vehroute', 'summary', 'tripinfo']
    #outputs = ['summary']

    path_sumo_outputs = f"../../../{path_sumo_outputs}/"
    print(path_sumo_outputs)

    for out in outputs:
        ET.SubElement(parent, f'{out}-output').set('value', os.path.join(f'{path_sumo_outputs}{out}.xml'))

        # Write xml
    output_dir = os.path.join(folders, 'sim.sumo.cfg')

    # End time
    hours_sec = hours * 3600;
    parent = tree.find('time')
    ET.SubElement(parent, 'end').set('value', f'{hours_sec}')

    tree.write(output_dir)
    return output_dir


def gen_DUArouter(net_file , trips, folders, i, hours):
    duarouter_conf = os.path.join('templates', 'duarouter.cfg.xml')  # duaroter.cfg file location
    #net_file = os.path.join(folders.parents_dir, 'templates', 'osm.net.xml')
    # Open original file
    tree = ET.parse(duarouter_conf)

    # Update trip input
    parent = tree.find('input')
    ET.SubElement(parent, 'net-file').set('value', f'{net_file}')
    ET.SubElement(parent, 'route-files').set('value', f"{trips}")

    # Update output
    parent = tree.find('output')
    output_name = os.path.join('rutas.rou.xml')
    ET.SubElement(parent, 'output-file').set('value', output_name)

    # Update seed number
    parent = tree.find('random_number')
    ET.SubElement(parent, 'seed').set('value', f'{i}')

    # End time
    hours_sec = hours * 3600;
    parent = tree.find('time')
    ET.SubElement(parent, 'end').set('value', f'{hours_sec}')

    # Write xml
    #original_path = os.path.dirname(trips)
    #cfg_name = os.path.join(original_path, f'{curr_name}_duarouter_{i}.cfg.xml')

    cfg_name = os.path.join(folders, 'dua.cfg.xml')

    tree.write(cfg_name)
    return cfg_name, output_name


def generate_edgeData(edge, out, time_difference):
    config_content = f"""<?xml version="1.0" encoding="UTF-8"?>
    <additional>
        <edgeData id="1" period="{time_difference}" type="emissions" file="{edge}" excludeEmpty="true" />
    </additional>
    """
    with open(out, 'w') as file:
        file.write(config_content)


if __name__ == '__main__':
    if len(sys.argv) != 10:
        print("Usage: python3 genOD_files.py <seed> <network.net.xml> <traffic_weights> <input_excel> <output_dir> <from_time> <to_time> <factor> <${base_name}/reportes_m>")
        sys.exit(1)

    seed = sys.argv[1]
    net_file = sys.argv[2]
    traffic_weights = sys.argv[3]
    input_excel = sys.argv[4]
    output_dir = sys.argv[5]
    from_time = sys.argv[6]
    to_time = sys.argv[7]
    factor = int(sys.argv[8])
    path_sumo_outputs = sys.argv[9]

    time_diff_in_file_name = time_diff(from_time, to_time)
    print(f"Generando {time_diff_in_file_name} horas de trafico ({time_diff_in_file_name} OD files)")
    df_xls_files = create_files_per_hour(traffic_weights, input_excel, output_dir, time_diff_in_file_name)
    list_OD_files = generate_output_file(df_xls_files, output_dir,1)
    od2_cfg_name, od2_result_file = gen_od2trips_cfg(list_OD_files, "districts.taz.xml", output_dir)

    # genera rou
    reroute_prob = 0.6
    trips_file = "od2.trip.xml"
    hours = float(to_time) - float(from_time)

    edgeData_output_xml = f"{output_dir}/edgeData_cfg.xml"
    edgeData_xml = f"../../../{path_sumo_outputs}/edgeData.xml"

    generate_edgeData(edgeData_xml, edgeData_output_xml, hours*3600)
    print("\nArchivo edgeData.xml generado correctamente")

    sumo_cfg = gen_routes(net_file , od2_cfg_name, trips_file, output_dir, seed, hours, reroute_prob, path_sumo_outputs, edgeData_xml)





